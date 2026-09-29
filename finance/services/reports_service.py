import pandas as pd

from django.utils import timezone
from django.db.models.functions import Coalesce

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from finance.models import Transaction, Budget
from finance.utils import get_month_range

class MonthlyReportService:

    def __init__(self, user, year=None, month=None):

        self.user = user
        reference_date = timezone.localtime()

        if year is not None:
            reference_date = reference_date.replace(year=year)

        if month is not None:
            reference_date = reference_date.replace(month=month)

        self.start, self.end = get_month_range(reference_date)

    def generate_excel(self):

        expenses = self.analyze_expense()
        incomes = self.analyze_income()
        budgets = self.get_budgets()
        accounts = self.get_accounts()

        budget_analysis = self.analyze_budget_vs_expenses(
            expenses,
            budgets,
        )

        expense_summary = (
            expenses
            .groupby(["currency", "category__name"])["amount"]
            .sum()
            .reset_index()
            .rename(
                columns={
                    "category__name": "category",
                    "amount": "total_gastado",
                }
            )
        )

        total_expenses = expenses["amount"].sum()
        total_incomes = incomes["amount"].sum()

        balance = total_incomes - total_expenses

        file_path = "monthly_report.xlsx"

        with pd.ExcelWriter(
            "monthly_report.xlsx",
            engine="openpyxl",
        ) as writer:

            self.create_summary_sheet(
                writer=writer,
                budget_analysis=budget_analysis,
                expense_summary=expense_summary,
                total_expenses=total_expenses,
                total_incomes=total_incomes,
                balance=balance,
            )

            budget_analysis.to_excel(
                writer,
                sheet_name="Presupuesto",
                index=False,
                startrow=2,
                startcol=1,
            )

            expense_summary.to_excel(
                writer,
                sheet_name="Gastos",
                index=False,
                startrow=2,
                startcol=1,
            )

            incomes.to_excel(
                writer,
                sheet_name="Ingresos",
                index=False,
                startrow=2,
                startcol=1,
            )

            accounts.to_excel(
                writer,
                sheet_name="Cuentas",
                index=False,
                startrow=2,
                startcol=1,
            )

            self.format_excel(writer)

            return file_path

    def analyze_expense(self):

        expenses = (
            Transaction.objects
            .filter(
                user=self.user,
                kind="expense",
            )
            .annotate(
                effective_date=Coalesce(
                    "occurred_at",
                    "created_at",
                )
            )
            .filter(
                effective_date__gte=self.start,
                effective_date__lt=self.end,
            )
            .values(
                "amount",
                "category__name",
                "account__name",
                "currency",
            )
        )

        df = pd.DataFrame.from_records(expenses)

        return df
    
    def analyze_income(self):

        incomes = ( 
            Transaction.objects.filter(
                user=self.user,
                kind= "income"
            )
            .annotate(
                effective_date=Coalesce(
                    "occurred_at",
                    "created_at",
                )
            )
            .filter(
                effective_date__gte=self.start,
                effective_date__lt=self.end,
            )
            .values(
                "amount",
                "category__name",
                "account__name",
                "currency",
            )
        )

        df = pd.DataFrame.from_records(incomes)

        return df
    
    def get_budgets(self):

        budgets = (
            Budget.objects.filter(
                user=self.user
            ).values(
                "category__name",
                "monthly_limit",
                "currency"
            )
        )

        df = pd.DataFrame.from_records(budgets)

        return df
    
    def get_accounts(self):

        accounts = (
            Transaction.objects
            .filter(
                user=self.user,
            )
            .annotate(
                effective_date=Coalesce(
                    "occurred_at",
                    "created_at",
                )
            )
            .filter(
                effective_date__gte=self.start,
                effective_date__lt=self.end,
            )
            .values(
                "account__name",
                "currency",
            )
        )

        df = pd.DataFrame.from_records(accounts)

        return df
    
    def analyze_budget_vs_expenses(self, expenses, budgets):

        expenses_by_category = (
            expenses
            .groupby(
                ["category__name","currency"]
            )["amount"]
            .sum()
            .reset_index()
        )

        expenses_by_category = expenses_by_category.rename(
            columns={
                "category__name": "category",
                "amount": "total_gastado",
            }
        )

        budgets = budgets.rename(
            columns={
                "category__name": "category",
                "monthly_limit": "limite_mensual",
            }
        )

        comparison = expenses_by_category.merge(
            budgets,
            on=["currency", "category"],
            how="left",
        )

        comparison["diferencia"] = (
            comparison["total_gastado"]
            - comparison["limite_mensual"]
        )

        return comparison[
            [
                "currency",
                "category",
                "total_gastado",
                "limite_mensual",
                "diferencia",
            ]
        ]
    
    def format_excel(self, writer):

        workbook = writer.book

        for worksheet in workbook.worksheets:

            # -------------------------
            # HOJA RESUMEN
            # -------------------------

            if worksheet.title == "Resumen":

                worksheet.freeze_panes = "B4"

                for row in worksheet.iter_rows():

                    for cell in row:

                        if isinstance(
                            cell.value,
                            (int, float),
                        ):

                            cell.number_format = "#,##0.00"

            # -------------------------
            # HOJAS DETALLADAS
            # -------------------------

            else:

                header_row = 3

                for cell in worksheet[header_row]:

                    if cell.value is None:
                        continue

                    cell.font = Font(
                        bold=True,
                        color="FFFFFF",
                    )

                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="1F4E78",
                    )

                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                    )

                worksheet.freeze_panes = "B4"

            # -------------------------
            # ANCHO DE COLUMNAS
            # -------------------------

            for column in worksheet.columns:

                max_length = 0

                column_letter = get_column_letter(
                    column[0].column
                )

                for cell in column:

                    if cell.value is not None:

                        max_length = max(
                            max_length,
                            len(str(cell.value)),
                        )

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 3,
                    35,
                )
                
    def create_summary_sheet(
        self,
        writer,
        budget_analysis,
        expense_summary,
        total_expenses,
        total_incomes,
        balance,
    ):

        workbook = writer.book

        worksheet = workbook.create_sheet("Resumen", 0)

        # -------------------------
        # TÍTULO
        # -------------------------

        worksheet.merge_cells("B1:F1")

        worksheet["B1"] = (
            f"REPORTE FINANCIERO — "
            f"{self.start.strftime('%B %Y').capitalize()}"
        )

        worksheet["B1"].font = Font(
            bold=True,
            size=18,
        )

        worksheet["B1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.row_dimensions[1].height = 30

        # -------------------------
        # RESUMEN DEL MES
        # -------------------------

        worksheet["B3"] = "RESUMEN DEL MES"

        worksheet["B3"].font = Font(
            bold=True,
            size=13,
        )

        worksheet["B5"] = "Ingresos"
        worksheet["C5"] = "Gastos"
        worksheet["D5"] = "Balance"

        worksheet["B6"] = total_incomes
        worksheet["C6"] = total_expenses
        worksheet["D6"] = balance

        for cell in worksheet[5]:

            if cell.column < 2 or cell.column > 4:
                continue

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

            cell.alignment = Alignment(
                horizontal="center",
            )

        for cell in worksheet[6][1:4]:

            cell.number_format = "#,##0.00"

            cell.font = Font(
                bold=True,
                size=12,
            )

        # -------------------------
        # PRESUPUESTO
        # -------------------------

        budget_start_row = 9

        worksheet.cell(
            row=budget_start_row,
            column=2,
            value="PRESUPUESTO",
        )

        worksheet.cell(
            row=budget_start_row,
            column=2,
        ).font = Font(
            bold=True,
            size=13,
        )

        headers = [
            "Currency",
            "Categoría",
            "Total gastado",
            "Límite mensual",
            "Diferencia",
            "Estado",
        ]

        header_row = budget_start_row + 2

        for column, value in enumerate(
            headers,
            start=2,
        ):

            cell = worksheet.cell(
                row=header_row,
                column=column,
                value=value,
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

            cell.alignment = Alignment(
                horizontal="center",
            )

        # -------------------------
        # DATOS DE PRESUPUESTO
        # -------------------------

        data_row = header_row + 1

        for _, row in budget_analysis.iterrows():

            worksheet.cell(
                row=data_row,
                column=2,
                value=row["currency"],
            )

            worksheet.cell(
                row=data_row,
                column=3,
                value=row["category"],
            )

            worksheet.cell(
                row=data_row,
                column=4,
                value=row["total_gastado"],
            )

            worksheet.cell(
                row=data_row,
                column=5,
                value=row["limite_mensual"],
            )

            worksheet.cell(
                row=data_row,
                column=6,
                value=row["diferencia"],
            )

            if pd.isna(row["limite_mensual"]):

                estado = "SIN LÍMITE"

            elif row["diferencia"] > 0:

                estado = "EXCEDIDO"

            elif row["diferencia"] == 0:

                estado = "LÍMITE"

            else:

                estado = "DENTRO DEL LÍMITE"

            worksheet.cell(
                row=data_row,
                column=7,
                value=estado,
            )

            data_row += 1

        # -------------------------
        # OBSERVACIONES
        # -------------------------

        observations_row = data_row + 2

        worksheet.cell(
            row=observations_row,
            column=2,
            value="OBSERVACIONES",
        )

        worksheet.cell(
            row=observations_row,
            column=2,
        ).font = Font(
            bold=True,
            size=13,
        )

        observation_row = observations_row + 2

        for _, row in budget_analysis.iterrows():

            category = row["category"]

            if pd.isna(row["limite_mensual"]):

                message = (
                    f"{category}: "
                    f"no tienes un límite mensual configurado."
                )

            elif row["diferencia"] > 0:

                message = (
                    f"{category}: "
                    f"superaste el límite mensual en "
                    f"{row['diferencia']:.2f} "
                    f"{row['currency']}."
                )

            elif row["diferencia"] == 0:

                message = (
                    f"{category}: "
                    f"alcanzaste exactamente el límite mensual."
                )

            else:

                message = (
                    f"{category}: "
                    f"estás dentro del límite mensual."
                )

            worksheet.cell(
                row=observation_row,
                column=2,
                value=message,
            )

            observation_row += 1

        # -------------------------
        # GASTOS POR CATEGORÍA
        # -------------------------

        expenses_row = observation_row + 2

        worksheet.cell(
            row=expenses_row,
            column=2,
            value="GASTOS POR CATEGORÍA",
        )

        worksheet.cell(
            row=expenses_row,
            column=2,
        ).font = Font(
            bold=True,
            size=13,
        )

        expense_header_row = expenses_row + 2

        for column, value in enumerate(
            ["Currency", "Categoría", "Total"],
            start=2,
        ):

            cell = worksheet.cell(
                row=expense_header_row,
                column=column,
                value=value,
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

        expense_data_row = expense_header_row + 1

        for _, row in expense_summary.iterrows():

            worksheet.cell(
                row=expense_data_row,
                column=2,
                value=row["currency"],
            )

            worksheet.cell(
                row=expense_data_row,
                column=3,
                value=row["category"],
            )

            worksheet.cell(
                row=expense_data_row,
                column=4,
                value=row["total_gastado"],
            )

            expense_data_row += 1